"""Export ranked candidates to Excel."""
from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models.schemas import RankedCandidate, CandidateProfile


def candidates_to_excel(candidates: List[CandidateProfile]) -> bytes:
    """Export raw candidate profiles (no ranking) to Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"

    headers = [
        "Filename", "Name", "Email", "Phone", "Location",
        "Total Experience (yrs)", "Current Company", "Notice Period",
        "Domain", "Skills", "Education", "Certifications", "Summary"
    ]
    _write_header(ws, headers)

    for row, c in enumerate(candidates, start=2):
        ws.cell(row=row, column=1, value=c.filename)
        ws.cell(row=row, column=2, value=c.name or "")
        ws.cell(row=row, column=3, value=c.email or "")
        ws.cell(row=row, column=4, value=c.phone or "")
        ws.cell(row=row, column=5, value=c.location or "")
        ws.cell(row=row, column=6, value=c.total_experience_years or 0)
        ws.cell(row=row, column=7, value=c.current_company or "")
        ws.cell(row=row, column=8, value=c.notice_period or "")
        ws.cell(row=row, column=9, value=c.domain or "")
        ws.cell(row=row, column=10, value=", ".join(c.skills))
        ws.cell(row=row, column=11, value="; ".join(c.education))
        ws.cell(row=row, column=12, value="; ".join(c.certifications))
        ws.cell(row=row, column=13, value=c.experience_summary or "")

    _auto_size(ws, headers)
    return _to_bytes(wb)


def ranked_to_excel(ranked: List[RankedCandidate]) -> bytes:
    """Export ranked candidate list with scores and breakdown."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranked Candidates"

    headers = [
        "Rank", "Name", "Filename", "Overall Score",
        "Skills Score", "Experience Score", "Education Score",
        "Domain Score", "Certifications Score", "Location/Notice Score",
        "Matched Skills", "Missing Skills", "Recommendation",
        "Email", "Phone", "Location", "Experience (yrs)",
        "Current Company", "Notice Period"
    ]
    _write_header(ws, headers)

    for row, r in enumerate(ranked, start=2):
        c = r.candidate
        b = r.breakdown
        ws.cell(row=row, column=1, value=r.rank)
        ws.cell(row=row, column=2, value=c.name or "")
        ws.cell(row=row, column=3, value=c.filename)
        ws.cell(row=row, column=4, value=r.overall_score)
        ws.cell(row=row, column=5, value=round(b.skills, 1))
        ws.cell(row=row, column=6, value=round(b.experience, 1))
        ws.cell(row=row, column=7, value=round(b.education, 1))
        ws.cell(row=row, column=8, value=round(b.domain, 1))
        ws.cell(row=row, column=9, value=round(b.certifications, 1))
        ws.cell(row=row, column=10, value=round(b.location_notice, 1))
        ws.cell(row=row, column=11, value=", ".join(r.matched_skills))
        ws.cell(row=row, column=12, value=", ".join(r.missing_skills))
        ws.cell(row=row, column=13, value=r.recommendation)
        ws.cell(row=row, column=14, value=c.email or "")
        ws.cell(row=row, column=15, value=c.phone or "")
        ws.cell(row=row, column=16, value=c.location or "")
        ws.cell(row=row, column=17, value=c.total_experience_years or 0)
        ws.cell(row=row, column=18, value=c.current_company or "")
        ws.cell(row=row, column=19, value=c.notice_period or "")

    _auto_size(ws, headers)
    _color_scores(ws, ranked)
    return _to_bytes(wb)


def _write_header(ws, headers):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24


def _auto_size(ws, headers):
    for col_idx, h in enumerate(headers, start=1):
        max_len = len(h)
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)), 50))
        ws.column_dimensions[col_letter].width = max_len + 2


def _color_scores(ws, ranked):
    """Color the overall score cell green/yellow/red based on value."""
    for row, r in enumerate(ranked, start=2):
        cell = ws.cell(row=row, column=4)
        if r.overall_score >= 75:
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
        elif r.overall_score >= 50:
            cell.fill = PatternFill("solid", fgColor="FFEB9C")
        else:
            cell.fill = PatternFill("solid", fgColor="FFC7CE")


def _to_bytes(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
